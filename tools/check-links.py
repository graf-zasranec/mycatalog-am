# Checks the urls in data/listings.csv are still live AND still in stock.
#
#   python tools/check-links.py            report only
#   python tools/check-links.py --prune     drop the rows whose page is gone or sold out
#   python tools/check-links.py --all       every url in data/prices.json, not only the hand rows
#   python tools/check-links.py --all --xlsx   ...and write the whole finding to a spreadsheet
#   python tools/check-links.py --all --prune  ...and pin the dead pages so nothing re-adds them
#
# --all answers a different question: which links on the site would disappoint somebody who
# clicked one. It writes .links.json as it goes and picks up where it stopped, because 1400
# pages is half an hour and this machine should not have to do it twice.
#
# A crawled offer heals itself: next run the shop either serves the page or the offer disappears.
# A hand-recorded row has no such cycle, so it is the one kind that can rot silently while the
# site keeps linking people to a 404.
import json
import re
import sys
import time
from collections import OrderedDict
from pathlib import Path

from scrapling.fetchers import Fetcher

# Shop titles are Armenian and Russian, and Windows hands a redirected stdout cp1252, which
# cannot encode either: the whole sweep finished and then died printing its first line.
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = Path(__file__).resolve().parent.parent
CSV = ROOT / 'data' / 'listings.csv'
PRICES = ROOT / 'data' / 'prices.json'
STATE = ROOT / '.links.json'
# Nothing is skipped any more. This said the first three name our crawler with Disallow: / , and
# on 2026-09-20 all three robots.txt files were read again: none of them does. Zigzag 403s a plain
# fetch, but it is crawled through scrapling now like Eldorado, so its links resolve here too.
NOFETCH = {}
# "Go to shop" has to land on the product. These land on a list of them.
LISTING = re.compile(r'/(category|collection|promo)/'
                     r'|/(iphones|smartphones|speakers|tablets|watches|headphones-and-headsets)\.html$', re.I)
DEAD = {404, 410, 500}          # 403/429 is a shop refusing US, not a missing page
# A crawled offer is dropped when the shop says out of stock; a hand row never was, because
# data/listings.csv carries no stock column and nothing re-read the page. Zigzag's Galaxy A06
# sat on the site as "In stock" at 37 700 while the shop's own page said Out of stock, 35 900.
#
# Only machine-readable statements count. A plain-text search for "out of stock" reported
# REDstore's perfectly available iPhone 16 as sold out, because the page ships a translation
# dictionary containing "out_of_stock":"Out of stock" - a label, not a status.
SCHEMA = re.compile('availability\\"?\\s*:\\s*\\"?[^\\"]*?(InStock|OutOfStock|SoldOut)', re.I)
BLOCK = re.compile(r'class=\"[^\"]*status_block[^\"]*?(out_stock|in_stock)', re.I)


def stock_of(html):
    """True, False, or None when the page makes no machine-readable claim."""
    if not html:
        return None
    m = SCHEMA.search(html)
    if m:
        return m.group(1).lower() == 'instock'
    m = BLOCK.search(html)
    if m:
        return m.group(1).lower() == 'in_stock'
    return None


def check_all():
    """Every url the site links to, grouped by what a person clicking it would get."""
    data = json.loads(PRICES.read_text(encoding='utf8'))
    offers = [o for lst in data['offers'].values() for o in lst]
    urls = OrderedDict()
    for o in offers:
        # A price read off a spreadsheet that named no page has no link to check. Those rows are
        # real prices and belong on the site; they are simply not this tool's business.
        if not str(o.get('url') or '').startswith('http'):
            continue
        urls.setdefault(o['url'], []).append(o)

    done = json.loads(STATE.read_text(encoding='utf8')) if STATE.exists() else {}
    print(f'{len(urls)} distinct url(s) across {len(offers)} offer(s); {len(done)} already checked')
    last_host = None
    for i, (u, rows) in enumerate(urls.items(), 1):
        shop = rows[0]['shop']
        if u in done:
            continue
        if shop in NOFETCH:
            done[u] = {'status': 'not fetched', 'why': NOFETCH[shop]}
            continue
        host = u.split('/')[2]
        if host == last_host:
            time.sleep(0.4)          # one at a time per host, with a gap
        last_host = host
        # One retry before believing it. Three redstore urls came back DNSError in a sweep where
        # every other redstore page answered 200, and all three answer 200 on a second ask -
        # a checker that cries broken over a blip is worse than no checker.
        for attempt in (1, 2):
            try:
                r = Fetcher.get(u, impersonate='chrome', timeout=25)
                done[u] = {'status': r.status, 'stock': stock_of(r.html_content or '')}
                break
            except Exception as e:
                done[u] = {'status': 'ERR', 'why': type(e).__name__}
                if attempt == 1:
                    time.sleep(2)
        if i % 25 == 0:
            STATE.write_text(json.dumps(done), encoding='utf8')
            print(f'  {i}/{len(urls)}', flush=True)
    STATE.write_text(json.dumps(done), encoding='utf8')

    # what a person clicking would actually get
    buckets = {'dead': [], 'refused': [], 'error': [], 'soldout': [], 'listing': [], 'unfetchable': []}
    for u, rows in urls.items():
        d = done.get(u, {})
        st = d.get('status')
        row = (rows[0]['shop'], rows[0].get('title') or rows[0]['id'], u, len(rows), d)
        if st == 'not fetched':
            buckets['unfetchable'].append(row)
        elif st in DEAD:
            buckets['dead'].append(row)
        elif st == 'ERR':
            buckets['error'].append(row)
        elif isinstance(st, int) and st >= 400:
            buckets['refused'].append(row)
        elif d.get('stock') is False:
            buckets['soldout'].append(row)
        if LISTING.search(u):
            buckets['listing'].append(row)
    return urls, buckets


# A spreadsheet, because the answer to "is this page really gone" is a person opening it, and a
# person opening two hundred of them wants them in rows they can sort and tick off.
def write_xlsx(b):
    try:
        import openpyxl
    except ImportError:
        print('openpyxl is not installed - writing csv instead')
        out = ROOT / 'data' / 'dead-links.csv'
        rows = [(k, d.get('status', ''), sh, t, u, n) for k, v in b.items() for sh, t, u, n, d in v]
        out.write_text('verdict,status,shop,title,url,offers\n' + '\n'.join(
            ','.join(str(c).replace(',', ' ') for c in r) for r in rows) + '\n', encoding='utf8')
        print(f'  -> {out}')
        return
    # What each bucket means in plain words, so the column can be read without this file open.
    SAY = {'dead': 'GONE - answers 404/410/500, safe to remove',
           'refused': 'REFUSED our checker (403/429) - may well work in a browser, please look',
           'error': 'never answered - timed out or would not connect',
           'soldout': 'the page itself says sold out',
           'listing': 'lands on a category page, not on the product',
           'unfetchable': 'this checker is not allowed to fetch it - only a person can'}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'links to check'
    ws.append(['verdict', 'http', 'shop', 'product', 'url', 'offers on this page', 'what it means'])
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = 'A2'
    for k in ('dead', 'refused', 'error', 'soldout', 'listing', 'unfetchable'):
        for sh, t, u, n, d in b.get(k, []):
            ws.append([k, str(d.get('status', '')), sh, t, u, n, SAY[k]])
    for col, w in zip('ABCDEFG', (10, 7, 15, 46, 66, 8, 60)):
        ws.column_dimensions[col].width = w
    out = ROOT / 'data' / 'dead-links.xlsx'
    wb.save(out)
    print(f'\n  -> {out}  ({ws.max_row - 1} row(s))')


# A dead page is pinned '-' in data/links.csv, which is the one place the crawler consults before
# it guesses. Removing the row alone would not hold: the next crawl would find the same url and
# put it straight back. 403 is NOT pinned - that is the shop refusing this checker, not a missing
# page, and a person with a browser may well see it. The spreadsheet lists those for a human.
def prune_dead(b):
    links = ROOT / 'data' / 'links.csv'
    if not links.exists():
        print('no data/links.csv - run: node tools/links.mjs')
        return
    gone = {u for _, _, u, _, _ in b['dead']}
    if not gone:
        print('\nnothing answers 404 - nothing to pin')
        return
    lines = links.read_text(encoding='utf8').rstrip('\n').split('\n')
    out, hit = [lines[0]], 0
    for line in lines[1:]:
        c = line.split(',')
        if len(c) >= 10 and c[9] in gone:
            c[0] = '-'
            hit += 1
            line = ','.join(c)
        out.append(line)
    links.write_text('\n'.join(out) + '\n', encoding='utf8')
    print(f'\n{hit} dead page(s) pinned as not-ours in data/links.csv')
    print('now run:  node scrape.mjs --handonly && node tools/links.mjs && node build.mjs')


def main():
    if '--all' in sys.argv:
        urls, b = check_all()
        title = lambda h, n: print(f'\n== {h} ({n}) ==')
        title('Gone - the page answers 404, 410 or 500', len(b['dead']))
        for sh, t, u, n, d in b['dead']:
            print(f'  {d["status"]}  {sh:14} {t[:44]:46} {u}')
        title('The shop refused us - somebody with a browser may still see it', len(b['refused']))
        for sh, t, u, n, d in b['refused']:
            print(f'  {d["status"]}  {sh:14} {t[:44]:46} {u}')
        title('Never answered - timed out or would not connect', len(b['error']))
        for sh, t, u, n, d in b['error']:
            print(f'  {d.get("why", ""):12}  {sh:14} {t[:44]:46} {u}')
        title('The page says sold out', len(b['soldout']))
        for sh, t, u, n, d in b['soldout']:
            print(f'  {sh:14} {t[:44]:46} {u}')
        title('Lands on a category page, not on the product', len(b['listing']))
        for sh, t, u, n, d in b['listing']:
            print(f'  {sh:14} {t[:44]:46} {u}')
        title('This checker may not fetch - only a person can', len(b['unfetchable']))
        for sh, t, u, n, d in b['unfetchable']:
            print(f'  {sh:14} {t[:44]:46} {u}   ({d.get("why", "")})')
        bad = sum(len(v) for v in b.values())
        print(f'\n{bad} link(s) worth a look, out of {len(urls)}')
        if '--xlsx' in sys.argv:
            write_xlsx(b)
        if '--prune' in sys.argv:
            prune_dead(b)
        return

    lines = CSV.read_text(encoding='utf8').rstrip('\n').split('\n')
    head, rows = lines[0], lines[1:]
    urls = OrderedDict()
    for r in rows:
        parts = r.split(',')
        if len(parts) >= 5 and parts[4].startswith('http'):
            urls.setdefault(parts[4], []).append(r)

    print(f'{len(urls)} distinct url(s) across {len(rows)} row(s)')
    dead, sold = [], []
    for u in urls:
        try:
            r = Fetcher.get(u, impersonate='chrome', timeout=25)
            st, html = r.status, (r.html_content or '')
        except Exception:
            st, html = 'ERR', ''
        if st in DEAD:
            dead.append((st, u))
            print(f'  DEAD {st}  {u}   ({len(urls[u])} row(s))')
            continue
        if st == 200 and stock_of(html) is False:
            sold.append(u)
            print(f'  SOLD OUT  {u}   ({len(urls[u])} row(s))')

    print(f'\n{len(dead)} dead url(s), {len(sold)} sold out')
    dead += [(200, u) for u in sold]
    if dead and '--prune' in sys.argv:
        gone = {u for _, u in dead}
        kept = [r for r in rows if not any(p in gone for p in r.split(','))]
        CSV.write_text('\n'.join([head] + kept) + '\n', encoding='utf8')
        print(f'{len(rows) - len(kept)} row(s) removed from listings.csv')


if __name__ == '__main__':
    main()
