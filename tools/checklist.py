# The hand-check list: one row per offer whose price, memory, RAM or colour needs a person's eye.
# Same rules as tools/todo-counts.mjs, but per offer, with the link and an empty column to answer in.
#   python tools/checklist.py [out.xlsx]
import json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

P = json.load(open('data/phones.json', encoding='utf-8'))
O = json.load(open('data/prices.json', encoding='utf-8'))['offers']
today = datetime.date.today()
out = sys.argv[1] if len(sys.argv) > 1 else f'C:/Users/hastv/Downloads/Better.am-checklist-{today}.xlsx'

def age(o):
    return (today - datetime.date.fromisoformat(o['seen'])).days if o.get('seen') else 999

rows = {'Price': [], 'Memory': [], 'RAM': [], 'Colour': []}
for p in sorted(P, key=lambda p: (p.get('category') or '', p.get('brand') or '', p.get('name') or '')):
    offs = O.get(p['id'], [])
    name = f"{p.get('brand', '')} {p.get('name', '')}".strip()
    stor = sorted({v['storage'] for v in p.get('variants') or [] if v.get('storage')})
    ram = sorted({v['ram'] for v in p.get('variants') or [] if v.get('ram')})
    cols = p.get('colors') or []
    low = {c.lower() for c in cols}
    for o in offs:
        def add(sheet, why, options=''):
            rows[sheet].append([why, p.get('category', ''), name, p['id'], o.get('shop', ''), o.get('price'),
                                o.get('storage'), o.get('ram'), o.get('color'), o.get('seen', ''),
                                o.get('title', ''), o.get('url', ''), options, '', ''])
        if o.get('price', 0) < 10000: add('Price', '1 under 10 000 AMD')
        elif age(o) > 7: add('Price', f'2 not re-read for {age(o)} days' if age(o) < 999 else '2 never re-read')
        elif len(offs) == 1: add('Price', '3 only shop - confirm price')
        if not p.get('variantUnit') and len(stor) > 1 and o.get('storage') is None:
            add('Memory', 'storage not stated', ' / '.join(map(str, stor)))
        if len(ram) > 1 and o.get('ram') is None:
            add('RAM', 'RAM not stated', ' / '.join(map(str, ram)))
        if len(cols) > 1 and not o.get('color'):
            add('Colour', 'colour not stated', ', '.join(cols))
        elif cols and o.get('color') and o['color'].lower() not in low:
            add('Colour', 'colour not in the product list', ', '.join(cols))

HEAD = ['What to check', 'Section', 'Product', 'Product id', 'Shop', 'Price AMD', 'Storage GB', 'RAM GB',
        'Colour', 'Last read', 'Shop title', 'Link', 'Product options', 'Correct value', 'Note']
WIDTH = [30, 11, 34, 30, 13, 11, 10, 8, 16, 11, 46, 14, 30, 16, 24]
wb = Workbook(); wb.remove(wb.active)
for sheet, rs in rows.items():
    rs.sort(key=lambda r: r[0])
    ws = wb.create_sheet(f'{sheet} ({len(rs)})')
    ws.append(HEAD)
    for r in rs:
        ws.append(r)
        c = ws.cell(ws.max_row, 12)
        if c.value: c.hyperlink, c.value, c.style = c.value, 'open', 'Hyperlink'
    for i, w in enumerate(WIDTH, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]: c.font = Font(bold=True)
    for c in ws['N'][1:] + ws['O'][1:]: c.fill = PatternFill('solid', fgColor='FFF4CC')
    for c in ws['F'][1:]: c.number_format = '#,##0'
    ws.freeze_panes, ws.auto_filter.ref = 'D2', ws.dimensions
wb.save(out)
print(out); [print(f'{len(v):5}  {k}') for k, v in rows.items()]
